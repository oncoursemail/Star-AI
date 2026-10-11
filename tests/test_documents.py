"""Synthetic document fixtures only; no model, desktop, network or user state."""
import io
import json
from pathlib import Path
import threading
import zipfile

import pytest

from star import documents
from star.documents import DocumentCancelled, DocumentError, DocumentStore


def text_file(tmp_path, name, text, encoding='utf-8'):
    path = tmp_path / name
    path.write_text(text, encoding=encoding, newline='')
    return path


def test_tail_retrieval_locations_and_hard_budget(tmp_path):
    lines = [f'Routine entry {index}: nothing distinctive here.' for index in range(1, 4001)]
    lines[-1] = 'ULTRAMARINE reconciliation result: invoice total 937.25.'
    path = text_file(tmp_path, 'long.txt', '\n'.join(lines))
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        manifest = store.ingest([path])[0]
        assert manifest['chunks'] == 4000
        context = store.context('What is the ULTRAMARINE invoice total?', 1600)
        assert '937.25' in context and 'line 4000' in context
        assert 'Coverage:' in context and 'omitted' in context
        assert 'D1-' in context and 'long.txt' in context
        assert store.context('ULTRAMARINE', 1600) == store.context('ULTRAMARINE', 1600)
        for budget in (0, 1, 20, 100, 500, 6000):
            assert len(store.context('ULTRAMARINE', budget)) <= budget
        cached = json.loads(store.db.execute('SELECT payload FROM document_payloads').fetchone()[0])
        assert cached['chunks'][-1]['text'] == lines[-1]


def test_restart_and_project_chat_isolation(tmp_path):
    home = tmp_path / 'home'
    path = text_file(tmp_path, 'source.csv', 'name,quantity\nAda,7\n')
    with DocumentStore(home, 'one', 'alpha') as store:
        store.ingest([path])
        identity = store.datasets()[0]['id']
    with DocumentStore(home, 'one', 'alpha') as store:
        assert store.datasets()[0]['id'] == identity
        assert 'Ada' in store.context('Ada')
    for project, chat in [('one', 'beta'), ('two', 'alpha'), ('one', None)]:
        with DocumentStore(home, project, chat) as store:
            assert store.context('Ada') == ''
            assert store.datasets() == []
    with DocumentStore(home, 'one', None) as store:
        store.ingest([path])
    with DocumentStore(home, 'one', 'legacy') as store:
        assert store.datasets() == []


def test_digest_change_replaces_reference_and_preserves_original(tmp_path):
    path = text_file(tmp_path, 'values.csv', 'value\n7\n')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        first = store.ingest([path])[0]
        first_id = store.datasets()[0]['id']
        path.write_text('value\n9\n', encoding='utf-8')
        second = store.ingest([path])[0]
        assert first['digest'] != second['digest']
        assert first_id != store.datasets()[0]['id']
        assert store.datasets()[0]['rows'] == [[9]]
        assert len(store.datasets()) == 1
        assert store.knowledge.artifact('project', first['digest']).read_text() == 'value\n7\n'


def test_derived_tamper_is_rejected_and_reattach_repairs(tmp_path):
    path = text_file(tmp_path, 'source.txt', 'Genuine source evidence.')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        store.db.execute("UPDATE document_payloads SET payload='{}'")
        store.db.commit()
        with pytest.raises(DocumentError, match='integrity'):
            store.context('evidence')
        store.ingest([path])
        assert 'Genuine source evidence' in store.context('evidence')


def test_immutable_tamper_is_rejected(tmp_path):
    path = text_file(tmp_path, 'source.txt', 'Genuine source evidence.')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        manifest = store.ingest([path])[0]
        artifact = store.knowledge.artifact('project', manifest['digest'])
        artifact.write_text('Altered evidence', encoding='utf-8')
        with pytest.raises(ValueError, match='digest|integrity|corrupt'):
            store.context('evidence')
        with pytest.raises(ValueError, match='digest|integrity|corrupt'):
            store.ingest([path])


def test_csv_multiline_utf8_bom_numeric_and_leading_zero(tmp_path):
    path = text_file(tmp_path, 'values.csv', 'name,amount,code\n"caf\u00e9\ncontinued",12.5,0012\n', 'utf-8-sig')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        dataset = store.datasets()[0]
        assert dataset['columns'] == ['name', 'amount', 'code']
        assert dataset['rows'] == [['caf\u00e9\ncontinued', '12.5', '0012']]
        assert dataset['row_sources'] == ['values.csv, lines 2-3']
        assert 'lines 2-3' in store.context('caf\u00e9')


def test_duplicate_long_headers_stay_bounded_and_can_be_calculated(tmp_path):
    from star.data_analysis import analyze
    header = 'x' * 200
    path = text_file(tmp_path, 'long-headers.csv', f'{header},{header}\n1,2\n')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        datasets = store.datasets()
        columns = datasets[0]['columns']
        assert len(columns) == len(set(columns)) == 2
        assert all(len(column) <= 200 for column in columns)
        result = analyze(datasets, {'dataset_id': datasets[0]['id'], 'operation': 'count'})
        assert result['value'] == result['input_rows'] == result['used_rows'] == 1


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16', 'utf-32'])
def test_text_boms(tmp_path, encoding):
    path = text_file(tmp_path, 'unicode.md', 'BOM sentinel caf\u00e9', encoding)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        assert 'BOM sentinel caf\u00e9' in store.context('sentinel')


def test_tsv_and_json_tables(tmp_path):
    tsv = text_file(tmp_path, 'table.tsv', 'label\tcount\nitem\t4\n')
    value = {'rows': [{'name': 'Ada', 'n': 3}, {'name': 'Bo', 'n': 5, 'meta': {'flag': True}}]}
    js = text_file(tmp_path, 'data.json', json.dumps(value))
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([tsv, js])
        datasets = {d['name'].split(' / ')[0]: d for d in store.datasets()}
        assert datasets['table.tsv']['rows'] == [['item', 4]]
        table = datasets['data.json']
        assert table['columns'] == ['name', 'n', 'meta']
        assert table['rows'] == [['Ada', 3, None], ['Bo', 5, '{"flag":true}']]
        assert 'data.json, $.rows[1]' in table['row_sources']
        assert any('Nested JSON' in warning for warning in table['warnings'])


def test_duplicate_content_filenames_have_distinct_dataset_ids(tmp_path):
    first = text_file(tmp_path, 'first.csv', 'value\n8\n')
    second = text_file(tmp_path, 'second.csv', 'value\n8\n')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([first, second])
        datasets = store.datasets()
        assert len({d['id'] for d in datasets}) == 2


def test_docx_paragraph_table_and_header_locations(tmp_path):
    docx = pytest.importorskip('docx')
    document = docx.Document()
    document.add_paragraph('First paragraph')
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('item', 'amount'), ('blue', '17')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph('Tail paragraph violet sentinel')
    document.sections[0].header.paragraphs[0].text = 'Header sentinel'
    path = tmp_path / 'source.docx'
    document.save(path)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        dataset = store.datasets()[0]
        assert dataset['columns'] == ['item', 'amount']
        assert dataset['rows'] == [['blue', '17']]
        assert dataset['row_sources'] == ['source.docx, table 1 row 2']
        assert 'paragraph 2' in store.context('violet')
        assert 'section 1 header paragraph 1' in store.context('Header sentinel')


def test_pptx_slide_and_speaker_notes_locations(tmp_path):
    pptx = pytest.importorskip('pptx')
    from pptx.util import Inches
    presentation = pptx.Presentation()
    for text in ['Opening slide', 'Last slide turquoise sentinel']:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
    presentation.slides[1].notes_slide.notes_text_frame.text = 'Speaker note evidence'
    path = tmp_path / 'source.pptx'
    presentation.save(path)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        assert 'slide 2 shape' in store.context('turquoise')
        assert 'slide 2 notes paragraph' in store.context('Speaker note')


def write_workbook(tmp_path):
    openpyxl = pytest.importorskip('openpyxl')
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Sales'
    sheet.append(['item', 'computed', 'uncached'])
    sheet.append(['blue', '=3+4', '=7+8'])
    path = tmp_path / 'source.xlsx'
    book.save(path)
    # Supply an explicitly stale cached result. No formula engine runs.
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<f>3+4</f><v></v>', b'<f>3+4</f><v>999</v>')
            target.writestr(member, data)
    return path


def test_xlsx_cached_formula_caveat_and_source_rows(tmp_path):
    path = write_workbook(tmp_path)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        dataset = store.datasets()[0]
        assert dataset['columns'] == ['item', 'computed', 'uncached']
        assert dataset['rows'] == [['blue', 999, None]]
        assert dataset['row_sources'] == ['source.xlsx, sheet Sales row 2']
        assert any('2 formula cells' in w and '1 cached values missing' in w for w in dataset['warnings'])
        context = store.context('computed', 3000)
        assert 'B2: cached formula (not recalculated): 999' in context
        assert '<missing cached value>' in context
        assert '=3+4' not in context


def write_pdf(path, text=None, encrypted=False):
    pypdf = pytest.importorskip('pypdf')
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text is not None:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted:
        writer.encrypt('synthetic-fixture')
    writer.write(path)
    return path


def test_pdf_page_provenance(tmp_path):
    path = write_pdf(tmp_path / 'source.pdf', 'PDF chartreuse sentinel')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        store.ingest([path])
        assert 'source.pdf, page 1' in store.context('chartreuse')


def test_pdf_limits_before_stream_expansion_and_configuration_restored(tmp_path, monkeypatch):
    pypdf = pytest.importorskip('pypdf')
    from pypdf.generic import NameObject, DecodedStreamObject
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(b'x' * 2000)
    page[NameObject('/Contents')] = writer._add_object(stream.flate_encode())
    path = tmp_path / 'large.pdf'
    writer.write(path)
    previous = pypdf.get_configuration()
    monkeypatch.setattr(documents, 'MAX_ZIP_MEMBER_BYTES', 1000)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        with pytest.raises(DocumentError, match='Malformed|safe content-stream'):
            store.ingest([path])
        assert store.context('anything') == ''
    assert pypdf.get_configuration() == previous


@pytest.mark.parametrize('variant', ['blank', 'encrypted', 'corrupt'])
def test_pdf_unreadable_rejected_honestly(tmp_path, variant):
    path = tmp_path / 'source.pdf'
    if variant == 'corrupt':
        path.write_bytes(b'%PDF-1.7\nnot a valid document')
    else:
        write_pdf(path, None if variant == 'blank' else 'encrypted evidence', variant == 'encrypted')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        with pytest.raises(DocumentError, match='extractable|Encrypted|Malformed'):
            store.ingest([path])
        assert store.context('anything') == ''


@pytest.mark.parametrize('name,text', [('ragged.csv', 'a,b\n1\n'), ('bad.json', '[NaN]'), ('broken.json', '{'), ('duplicate.json', '{"a":1,"a":2}'), ('binary.txt', 'hi\x00there')])
def test_malformed_text_rejected(tmp_path, name, text):
    path = text_file(tmp_path, name, text)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        with pytest.raises(DocumentError):
            store.ingest([path])
        assert store.datasets() == []


def test_extraction_file_cell_and_zip_limits(tmp_path, monkeypatch):
    home = tmp_path / 'home'
    path = text_file(tmp_path, 'long.txt', 'x' * 101)
    with DocumentStore(home, 'project', 'chat') as store:
        monkeypatch.setattr(documents, 'MAX_EXTRACTED_CHARS', 100)
        with pytest.raises(DocumentError, match='character'):
            store.ingest([path])
        monkeypatch.setattr(documents, 'MAX_EXTRACTED_CHARS', 1_000_000)
        monkeypatch.setattr(documents, 'MAX_FILE_BYTES', 50)
        with pytest.raises(DocumentError, match='larger|file-size'):
            store.ingest([path])
        monkeypatch.setattr(documents, 'MAX_FILE_BYTES', 32 * 1024 * 1024)
        monkeypatch.setattr(documents, 'MAX_CELLS', 5)
        csv = text_file(tmp_path, 'many.csv', 'a,b\n1,2\n3,4\n')
        with pytest.raises(DocumentError, match='cell'):
            store.ingest([csv])
        archive = tmp_path / 'bomb.docx'
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('large.xml', b'a' * 1_100_000)
        with pytest.raises(DocumentError, match='expansion'):
            store.ingest([archive])
        assert store.context('anything') == ''


def test_macro_and_archive_traversal_rejected(tmp_path):
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        for member, reason in [('xl/vbaProject.bin', 'Macro'), ('../escape.xml', 'unsafe')]:
            path = tmp_path / 'unsafe.xlsx'
            with zipfile.ZipFile(path, 'w') as archive:
                archive.writestr(member, b'fixture')
            with pytest.raises(DocumentError, match=reason):
                store.ingest([path])


def test_missing_optional_dependency_has_clear_error(tmp_path, monkeypatch):
    path = write_pdf(tmp_path / 'source.pdf', 'Evidence')
    original = documents.importlib.import_module
    def missing(name):
        if name == 'pypdf':
            raise ImportError('synthetic missing dependency')
        return original(name)
    monkeypatch.setattr(documents.importlib, 'import_module', missing)
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        with pytest.raises(DocumentError, match='optional document dependency pypdf'):
            store.ingest([path])


def test_cancellation_and_batch_failure_leave_no_references(tmp_path):
    good = text_file(tmp_path, 'good.txt', 'Good evidence')
    bad = text_file(tmp_path, 'bad.csv', 'a,b\n1\n')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        event = threading.Event()
        event.set()
        with pytest.raises(DocumentCancelled):
            store.ingest([good], event)
        with pytest.raises(DocumentError):
            store.ingest([good, bad])
        assert store.context('Good') == ''
        class AfterUnits:
            count = 0
            def is_set(self):
                self.count += 1
                return self.count > 12
        many = text_file(tmp_path, 'many.txt', '\n'.join('evidence' for _ in range(100)))
        with pytest.raises(DocumentCancelled):
            store.ingest([many], AfterUnits())
        assert store.context('evidence') == ''


def test_invalid_scopes_and_budget(tmp_path):
    with pytest.raises(ValueError):
        DocumentStore(tmp_path / 'home', '../invalid', 'chat')
    with pytest.raises(DocumentError):
        DocumentStore(tmp_path / 'home', 'project', '')
    with DocumentStore(tmp_path / 'home', 'project', 'chat') as store:
        for budget in (-1, 2.5, True):
            with pytest.raises(DocumentError):
                store.context('anything', budget)
        store.close()
        store.close()
