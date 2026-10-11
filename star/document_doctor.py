"""Exercise packaged attachment readers with authored, disposable sources.

No application profile, engine, provider or network is used. This runs only for
the executable's doctor commands, so missing parser resources fail packaging.
"""
from pathlib import Path
from tempfile import TemporaryDirectory

from .data_analysis import analyze
from .documents import DocumentStore


def check_document_bundle() -> tuple[str, ...]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    with TemporaryDirectory(prefix="star-document-doctor-") as directory:
        root = Path(directory)
        sources = root / "sources"
        sources.mkdir()
        markers = {extension: "PACKAGED_" + extension.upper()
                   for extension in ("txt", "pdf", "docx", "pptx", "csv", "tsv", "xlsx", "json")}
        (sources / "sample.txt").write_text(markers["txt"], encoding="utf-8")
        word = Document()
        word.add_paragraph(markers["docx"])
        word.save(sources / "sample.docx")
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = markers["pptx"]
        slides.save(sources / "sample.pptx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["marker", "amount"])
        sheet.append([markers["xlsx"], 0.1])
        sheet.append(["second", 0.2])
        workbook.save(sources / "sample.xlsx")
        workbook.close()
        for extension, separator in (("csv", ","), ("tsv", "\t")):
            (sources / ("sample." + extension)).write_text(
                f"marker{separator}amount\n{markers[extension]}{separator}0.1\nsecond{separator}0.2\n",
                encoding="utf-8")
        (sources / "sample.json").write_text(
            '[{"marker":"PACKAGED_JSON","amount":0.1},{"marker":"second","amount":0.2}]',
            encoding="utf-8")
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (PACKAGED_PDF) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(sources / "sample.pdf")
        writer.close()

        with DocumentStore(root / "profile", "packaged-readers", "doctor") as store:
            store.ingest([sources / ("sample." + extension) for extension in markers])
            for extension, marker in markers.items():
                evidence = store.context(marker, budget_chars=8000)
                if marker not in evidence or "sample." + extension not in evidence:
                    raise RuntimeError(f"Packaged {extension.upper()} evidence extraction failed")
            datasets = store.datasets()
            for extension in ("csv", "tsv", "xlsx", "json"):
                dataset = next(item for item in datasets if item["name"].startswith("sample." + extension))
                result = analyze([dataset], {"dataset_id": dataset["id"], "operation": "sum", "column": "amount"})
                if result["value"] != "0.3" or not result["exact"] or result["used_rows"] != 2:
                    raise RuntimeError(f"Packaged {extension.upper()} calculation failed")
        return tuple(markers)
