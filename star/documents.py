"""Bounded, source-located document evidence; no code, macros or formulas run.

Originals live in KnowledgeStore's immutable artifact store. Derived extraction
and references are separate from trusted lessons and scoped to a project/chat.
Retrieval is deterministic and never calls a model or an external service.
"""
from __future__ import annotations

from collections import Counter
import csv
from datetime import date, datetime, time
import hashlib
import importlib
import io
import json
import math
from pathlib import Path
import re
import sqlite3
import zipfile

from .knowledge import KnowledgeStore, _project
from . import rbx_xml

# Binary Roblox extensions are routed here only to give an actionable refusal.
DOCUMENT_EXTENSIONS = (frozenset({'.txt', '.md', '.pdf', '.docx', '.pptx', '.csv', '.tsv', '.xlsx', '.json'})
                       | rbx_xml.XML_EXTENSIONS | rbx_xml.BINARY_EXTENSIONS)
MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_EXTRACTED_CHARS = 1_000_000
MAX_CELLS = 100_000
MAX_COLUMNS = 1_000
MAX_UNZIP_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_MEMBERS = 10_000
MAX_PAGES = 2_000
MAX_DOCUMENTS = 32
MAX_PAYLOAD_BYTES = 12 * 1024 * 1024
CHUNK_CHARS = 1400
EXTRACTION_VERSION = 1
_MIMES = {'.txt': 'text/plain', '.md': 'text/markdown', '.pdf': 'application/pdf',
          '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
          '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
          '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
          '.csv': 'text/csv', '.tsv': 'text/tab-separated-values', '.json': 'application/json',
          '.rbxmx': 'application/xml', '.rbxlx': 'application/xml'}
_WORDS = re.compile(r"\w+", re.UNICODE)
_STOP = frozenset('the a an of to and or in on for is are was were what how which please document file tell me about'.split())
_NUMBER = re.compile(r'-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?\Z')


class DocumentError(ValueError):
    """A document could not be ingested or its saved evidence is invalid."""


class DocumentCancelled(RuntimeError):
    pass


def _check(cancel_event=None):
    if cancel_event is not None and cancel_event.is_set():
        raise DocumentCancelled('Document processing cancelled')


def _dependency(module: str, package: str):
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        raise DocumentError(f'Install the optional document dependency {package} to read this format') from exc


def _label(value) -> str:
    # Source labels cannot create additional lines or app context delimiters.
    return ' '.join(str(value).replace('[STAR', '(STAR').replace('[END STAR', '(END STAR').split())[:200]


def _decode(raw: bytes) -> str:
    try:
        if raw.startswith((b'\xff\xfe\x00\x00', b'\x00\x00\xfe\xff')):
            text = raw.decode('utf-32')
        elif raw.startswith((b'\xff\xfe', b'\xfe\xff')):
            text = raw.decode('utf-16')
        else:
            text = raw.decode('utf-8-sig')
    except UnicodeError as exc:
        raise DocumentError('Text must be UTF-8 or Unicode with a byte-order mark') from exc
    if '\x00' in text:
        raise DocumentError('Binary data is not a supported text document')
    if len(text) > MAX_EXTRACTED_CHARS:
        raise DocumentError(f'Document exceeds the {MAX_EXTRACTED_CHARS:,}-character extraction limit')
    return text


def _scalar(value, *, numeric_text=False):
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DocumentError('Non-finite numeric values are not supported')
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        if numeric_text and _NUMBER.fullmatch(value.strip()):
            raw = value.strip()
            try:
                # Keep decimal/exponent lexemes exact for the Decimal calculator;
                # converting them through a binary float can destroy digits.
                return raw if any(c in raw for c in '.eE') else int(raw)
            except ValueError:
                pass
        return value
    raise DocumentError('A data cell contains a non-scalar value')


def _columns(values) -> list[str]:
    result, used = [], Counter()
    for index, value in enumerate(values, 1):
        base = _label(value) if value is not None and str(value).strip() else f'column_{index}'
        used[base] += 1
        suffix = '' if used[base] == 1 else f'_{used[base]}'
        name = base[:200 - len(suffix)] + suffix
        while name in result:
            used[base] += 1
            suffix = f'_{used[base]}'
            name = base[:200 - len(suffix)] + suffix
        result.append(name)
    return result


class _Evidence:
    def __init__(self, cancel_event=None):
        self.cancel = cancel_event
        self.chunks, self.tables, self.warnings = [], [], []
        self.chars = self.cells = 0

    def add(self, location: str, text: str):
        _check(self.cancel)
        text = str(text).strip()
        self.chars += len(text)
        if self.chars > MAX_EXTRACTED_CHARS:
            raise DocumentError(f'Document exceeds the {MAX_EXTRACTED_CHARS:,}-character extraction limit')
        for offset in range(0, len(text), CHUNK_CHARS):
            _check(self.cancel)
            part = text[offset:offset + CHUNK_CHARS]
            label = location if len(text) <= CHUNK_CHARS else f'{location}, chars {offset + 1}-{offset + len(part)}'
            self.chunks.append({'location': _label(label), 'text': part})

    def table(self, label: str, columns: list[str], rows: list[list], locations: list[str], warnings=()):
        if len(columns) > MAX_COLUMNS:
            raise DocumentError(f'Table exceeds the {MAX_COLUMNS:,}-column limit')
        self.cells += len(columns) * (len(rows) + 1)
        if self.cells > MAX_CELLS:
            raise DocumentError(f'Document exceeds the {MAX_CELLS:,}-cell limit')
        self.tables.append({'label': _label(label), 'columns': columns, 'rows': rows,
                            'locations': locations, 'warnings': list(warnings)})

    def result(self):
        if not self.chunks:
            raise DocumentError('No extractable text; scanned or image-only documents need OCR, which is not supported')
        return {'chunks': self.chunks, 'tables': self.tables, 'warnings': self.warnings, 'chars': self.chars}


def _text(text: str, out: _Evidence):
    for number, line in enumerate(text.splitlines(), 1):
        out.add(f'line {number}', line)


def _delimited(text: str, out: _Evidence, suffix: str):
    reader = csv.reader(io.StringIO(text, newline=''), delimiter='\t' if suffix == '.tsv' else ',', strict=True)
    rows, locations, header = [], [], None
    previous_line = 0
    for row in reader:
        _check(out.cancel)
        first, last = previous_line + 1, reader.line_num
        previous_line = last
        if not row:
            continue
        if len(row) > MAX_COLUMNS:
            raise DocumentError('Data row exceeds the column limit')
        location = f'line {first}' if first == last else f'lines {first}-{last}'
        out.add(location, ' | '.join(row))
        if header is None:
            header = _columns(row)
        else:
            if len(row) != len(header):
                raise DocumentError(f'CSV/TSV row at {location} has {len(row)} cells; expected {len(header)}')
            if (len(rows) + 2) * len(header) > MAX_CELLS:
                raise DocumentError('Data file exceeds the cell limit')
            rows.append([_scalar(value, numeric_text=True) for value in row])
            locations.append(location)
    if header:
        out.table('table 1', header, rows, locations)


def _json(text: str, out: _Evidence):
    def invalid_constant(value):
        raise DocumentError(f'Invalid JSON numeric constant: {value}')
    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise DocumentError(f'Duplicate JSON key: {_label(key)}')
            result[key] = value
        return result
    value = json.loads(text, parse_float=str, parse_constant=invalid_constant, object_pairs_hook=unique_object)
    _text(text, out)
    candidates = [('$', value)] if isinstance(value, list) else []
    if isinstance(value, dict):
        candidates = [(f'$.{key}', item) for key, item in value.items() if isinstance(item, list)]
        if not candidates and value and all(not isinstance(item, (dict, list)) for item in value.values()):
            candidates = [('$', [value])]
    if len(candidates) > 32:
        raise DocumentError('JSON exceeds the 32-table limit')
    for path, array in candidates:
        _check(out.cancel)
        if not array:
            continue
        warnings = []
        if all(isinstance(row, dict) for row in array):
            keys = list(dict.fromkeys(key for row in array for key in row))
            columns = _columns(keys)
            raw_rows = [[row.get(key) for key in keys] for row in array]
        elif all(isinstance(row, list) for row in array):
            width = len(array[0])
            if any(len(row) != width for row in array):
                raise DocumentError('JSON table rows have inconsistent widths')
            columns, raw_rows = _columns([None] * width), array
        elif all(not isinstance(row, (dict, list)) for row in array):
            columns, raw_rows = ['value'], [[row] for row in array]
        else:
            out.warnings.append(f'{_label(path)} is not a rectangular table; its source text remains available')
            continue
        if len(columns) > MAX_COLUMNS or len(columns) * (len(raw_rows) + 1) > MAX_CELLS:
            raise DocumentError('JSON table exceeds the cell or column limit')
        rows = []
        for row in raw_rows:
            _check(out.cancel)
            values = []
            for item in row:
                if isinstance(item, (dict, list)):
                    item = json.dumps(item, ensure_ascii=False, separators=(',', ':'))
                    if not warnings:
                        warnings.append('Nested JSON cells are retained as text, not numeric values')
                values.append(_scalar(item))
            rows.append(values)
        out.table(path, columns, rows, [f'{path}[{index}]' for index in range(len(rows))], warnings)


def _safe_zip(path: Path, cancel_event=None):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise DocumentError('Office archive contains too many members')
        total = 0
        for member in members:
            _check(cancel_event)
            name = member.filename.replace('\\', '/')
            if name.startswith('/') or '..' in name.split('/') or member.flag_bits & 1:
                raise DocumentError('Office archive has unsafe paths or encrypted members')
            if name.lower().endswith('vbaproject.bin'):
                raise DocumentError('Macro-bearing Office files are not supported')
            total += member.file_size
            if (member.file_size > MAX_ZIP_MEMBER_BYTES or total > MAX_UNZIP_BYTES
                    or member.file_size > 1_000_000 and member.file_size > max(1, member.compress_size) * 500):
                raise DocumentError('Office archive exceeds safe expansion limits')


def _docx(path: Path, out: _Evidence):
    docx = _dependency('docx', 'python-docx')
    from docx.table import Table
    document = docx.Document(path)
    paragraph, table_number = 0, 0
    for block in document.iter_inner_content():
        _check(out.cancel)
        if isinstance(block, Table):
            table_number += 1
            values, locations, cell_count = [], [], 0
            for number, row in enumerate(block.rows, 1):
                _check(out.cancel)
                cells = [cell.text for cell in row.cells]
                cell_count += len(cells)
                if cell_count + out.cells > MAX_CELLS or len(cells) > MAX_COLUMNS:
                    raise DocumentError('Word table exceeds the cell limit')
                location = f'table {table_number} row {number}'
                out.add(location, ' | '.join(cells))
                values.append(cells)
                locations.append(location)
            if values:
                width = max(map(len, values))
                values = [row + [None] * (width - len(row)) for row in values]
                out.table(f'table {table_number}', _columns(values[0]), values[1:], locations[1:])
        else:
            paragraph += 1
            out.add(f'paragraph {paragraph}', block.text)
    for number, section in enumerate(document.sections, 1):
        for label, part in (('header', section.header), ('footer', section.footer)):
            for index, item in enumerate(part.paragraphs, 1):
                out.add(f'section {number} {label} paragraph {index}', item.text)
    out.warnings.append('Word body paragraphs and tables plus header/footer paragraphs extracted; nested/header/footer tables, embedded objects and drawing text are not interpreted')


def _pptx(path: Path, out: _Evidence):
    pptx = _dependency('pptx', 'python-pptx')
    presentation = pptx.Presentation(path)
    if len(presentation.slides) > MAX_PAGES:
        raise DocumentError('Presentation exceeds the slide limit')
    for number, slide in enumerate(presentation.slides, 1):
        start = len(out.chunks)
        def shapes(items):
            for index, shape in enumerate(items, 1):
                _check(out.cancel)
                if hasattr(shape, 'shapes'):
                    shapes(shape.shapes)
                if shape.has_text_frame:
                    for paragraph, text in enumerate(shape.text_frame.paragraphs, 1):
                        out.add(f'slide {number} shape {shape.shape_id} paragraph {paragraph}', text.text)
                if shape.has_table:
                    values, locations, cell_count = [], [], 0
                    for row_number, row in enumerate(shape.table.rows, 1):
                        _check(out.cancel)
                        cells = [cell.text for cell in row.cells]
                        cell_count += len(cells)
                        if cell_count + out.cells > MAX_CELLS or len(cells) > MAX_COLUMNS:
                            raise DocumentError('Slide table exceeds the cell or column limit')
                        values.append(cells)
                        location = f'slide {number} table {shape.shape_id} row {row_number}'
                        locations.append(location)
                        out.add(location, ' | '.join(values[-1]))
                    if values:
                        out.table(f'slide {number} table {shape.shape_id}', _columns(values[0]), values[1:], locations[1:])
        shapes(slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                for index, paragraph in enumerate(frame.paragraphs, 1):
                    out.add(f'slide {number} notes paragraph {index}', paragraph.text)
        if start == len(out.chunks):
            out.warnings.append(f'Slide {number} has no extractable text; pictures are not OCRed')
    out.warnings.append('Slide text, tables and existing speaker notes extracted; images, charts and embedded objects are not interpreted')


def _xlsx(path: Path, out: _Evidence):
    openpyxl = _dependency('openpyxl', 'openpyxl')
    # Artifact filenames are digests, so pass streams to avoid suffix-based
    # rejection. These handles remain open while read-only worksheets stream.
    formula_stream, cache_stream = path.open('rb'), path.open('rb')
    formulas = cached = None
    try:
        formulas = openpyxl.load_workbook(formula_stream, read_only=True, data_only=False, keep_links=False, keep_vba=False)
        cached = openpyxl.load_workbook(cache_stream, read_only=True, data_only=True, keep_links=False, keep_vba=False)
        for sheet in formulas.worksheets:
            _check(out.cancel)
            if sheet.max_column is None or sheet.max_row is None:
                raise DocumentError('Worksheet dimensions are missing; save the workbook with a valid used range')
            if sheet.max_column > MAX_COLUMNS or sheet.max_row * sheet.max_column > MAX_CELLS:
                raise DocumentError('Worksheet used grid exceeds the cell or column limit')
            values, locations, formula_count, missing = [], [], 0, 0
            width = max(1, sheet.max_column or 1)
            known = cached[sheet.title]
            for row_number, (row, cached_row) in enumerate(zip(sheet.iter_rows(max_col=width), known.iter_rows(max_col=width)), 1):
                _check(out.cancel)
                if (len(values) + 1) * width + out.cells > MAX_CELLS:
                    raise DocumentError('Workbook exceeds the cell limit')
                current, display = [], []
                for cell, saved in zip(row, cached_row):
                    if cell.data_type == 'f':
                        formula_count += 1
                        value = _scalar(saved.value)
                        missing += value is None
                        shown = f'{cell.coordinate}: cached formula (not recalculated): {value if value is not None else "<missing cached value>"}'
                    else:
                        value = _scalar(cell.value)
                        shown = f'{cell.coordinate}: {value}'
                    current.append(value)
                    display.append(shown)
                if any(value is not None for value in current) or any(cell.data_type == 'f' for cell in row):
                    location = f'sheet {_label(sheet.title)} row {row_number}'
                    out.add(location, ' | '.join(display))
                    values.append(current)
                    locations.append(location)
            if values:
                warnings = []
                if formula_count:
                    warnings.append(f'{formula_count} formula cells use cached values only; formulas were not recalculated; {missing} cached values missing')
                out.table(f'sheet {_label(sheet.title)}', _columns(values[0]), values[1:], locations[1:], warnings)
        out.warnings.append('Excel values and cached formula results only; macros, formulas and external links are never executed or refreshed')
    finally:
        if formulas is not None:
            formulas.close()
        if cached is not None:
            cached.close()
        formula_stream.close()
        cache_stream.close()


def _pdf(path: Path, out: _Evidence):
    pypdf = _dependency('pypdf', 'pypdf')
    if not hasattr(pypdf, 'apply_configuration'):
        raise DocumentError('Install pypdf>=6.20 for bounded PDF parsing')
    # Resource checks must apply before decompression and page-tree traversal.
    # These settings are local to this execution context, not process globals.
    with pypdf.apply_configuration(
            maximum_declared_stream_length=MAX_ZIP_MEMBER_BYTES,
            array_based_stream_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            zlib_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            lzw_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            run_length_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            jbig2_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            brotli_maximum_output_length=MAX_ZIP_MEMBER_BYTES,
            image_maximum_buffer_size=MAX_ZIP_MEMBER_BYTES,
            stream_decoding_work_maximum_length=MAX_UNZIP_BYTES,
            page_tree_maximum_entries=MAX_PAGES * 2,
            page_tree_maximum_depth=64,
            xform_maximum_invocations_per_extraction=500,
            jbig2dec_binary=None):
        reader = pypdf.PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise DocumentError('Encrypted PDFs are not supported; provide an unencrypted copy')
        if len(reader.pages) > MAX_PAGES:
            raise DocumentError('PDF exceeds the page limit')
        for number, page in enumerate(reader.pages, 1):
            _check(out.cancel)
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > MAX_ZIP_MEMBER_BYTES:
                raise DocumentError(f'PDF page {number} exceeds the safe content-stream limit')
            text = page.extract_text() or ''
            if text.strip():
                out.add(f'page {number}', text)
            else:
                out.warnings.append(f'Page {number} has no extractable text; scanned content is not OCRed')
    out.warnings.append('PDF text extracted by page; layout, pictures and table structure are not guaranteed')


def _roblox(path: Path, suffix: str, out: _Evidence):
    try:
        model = rbx_xml.read(path, suffix=suffix, check=lambda: _check(out.cancel))
    except rbx_xml.RbxXMLError as exc:
        raise DocumentError(str(exc)) from exc
    for location, text, kind in model.evidence():
        first = len(out.chunks)
        out.add(location, text)
        for chunk in out.chunks[first:]:
            chunk['kind'] = kind
    out.warnings.extend(model.warnings)
    out.warnings.append('Roblox XML outline and inline Luau only; lexical script summaries do not resolve dynamic calls. Assets and scripts are never fetched or executed.')


def _extract(path: Path, suffix: str, cancel_event=None) -> dict:
    out = _Evidence(cancel_event)
    _check(cancel_event)
    try:
        if suffix in {'.docx', '.pptx', '.xlsx'}:
            _safe_zip(path, cancel_event)
        if suffix in {'.txt', '.md', '.csv', '.tsv', '.json'}:
            with path.open('rb') as stream:
                raw = stream.read(MAX_FILE_BYTES + 1)
            if len(raw) > MAX_FILE_BYTES:
                raise DocumentError('Document exceeds the file-size limit')
            text = _decode(raw)
            if suffix in {'.csv', '.tsv'}:
                _delimited(text, out, suffix)
            elif suffix == '.json':
                _json(text, out)
            else:
                _text(text, out)
        elif suffix == '.pdf':
            _pdf(path, out)
        elif suffix == '.docx':
            _docx(path, out)
        elif suffix == '.pptx':
            _pptx(path, out)
        elif suffix == '.xlsx':
            _xlsx(path, out)
        elif suffix in rbx_xml.XML_EXTENSIONS | rbx_xml.BINARY_EXTENSIONS:
            _roblox(path, suffix, out)
        return out.result()
    except (DocumentError, DocumentCancelled):
        raise
    except Exception as exc:
        raise DocumentError(f'Malformed or unsupported {suffix} document ({type(exc).__name__})') from exc


def _hash(value: str) -> str:
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


class DocumentStore:
    def __init__(self, home: Path, project: str, chat_id: str | None):
        self.home = Path(home)
        self.project = _project(project)
        chat = None if chat_id is None else str(chat_id)
        if chat is not None and (not chat or len(chat) > 200 or any(ord(c) < 32 for c in chat)):
            raise DocumentError('Invalid document chat scope')
        self.chat = 'none:' if chat is None else 'chat:' + chat
        self.home.mkdir(parents=True, exist_ok=True)
        self.knowledge = KnowledgeStore(self.home / 'knowledge')
        self.db = sqlite3.connect(self.home / 'documents.sqlite3', timeout=10)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.executescript('''
          CREATE TABLE IF NOT EXISTS document_payloads (
            digest TEXT NOT NULL, extension TEXT NOT NULL, version INTEGER NOT NULL,
            payload TEXT NOT NULL, payload_sha TEXT NOT NULL,
            PRIMARY KEY(digest,extension,version));
          CREATE TABLE IF NOT EXISTS document_refs (
            project TEXT NOT NULL, chat TEXT NOT NULL, source_key TEXT NOT NULL,
            digest TEXT NOT NULL, extension TEXT NOT NULL, name TEXT NOT NULL,
            PRIMARY KEY(project,chat,source_key));
        ''')

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()

    def close(self):
        if self.db is not None:
            self.db.close()
            self.db = None
        if self.knowledge is not None:
            self.knowledge.close()
            self.knowledge = None

    def ingest(self, paths: list[Path], cancel_event=None) -> list[dict]:
        if not isinstance(paths, list) or len(paths) > 8:
            raise DocumentError('Choose at most eight document attachments at a time')
        prepared = []
        for supplied in paths:
            _check(cancel_event)
            original = Path(supplied).expanduser()
            if original.is_symlink():
                raise DocumentError('Symlink document attachments are not supported')
            path = original.resolve(strict=True)
            suffix = path.suffix.lower()
            if suffix not in DOCUMENT_EXTENSIONS:
                raise DocumentError(f'Unsupported document extension: {suffix or "none"}')
            if suffix in rbx_xml.BINARY_EXTENSIONS:
                raise DocumentError(rbx_xml.BINARY_MESSAGE)
            if suffix in rbx_xml.XML_EXTENSIONS and path.stat().st_size > rbx_xml.MAX_FILE_BYTES:
                raise DocumentError('Roblox XML exceeds the file-size limit')
            if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
                raise DocumentError('Document is missing or larger than 32 MiB')
            key = _hash(str(path))
            digest, immutable = self.knowledge.ingest_artifact(
                project=self.project, path=path, origin='document:' + key,
                media_type=_MIMES[suffix])
            self._verify_original(digest, path.name)
            cached = self.db.execute('SELECT payload,payload_sha FROM document_payloads WHERE digest=? AND extension=? AND version=?',
                                     (digest, suffix, EXTRACTION_VERSION)).fetchone()
            if cached is not None and _hash(cached['payload']) == cached['payload_sha']:
                payload = cached['payload']
                evidence = json.loads(payload)
            else:
                evidence = _extract(immutable, suffix, cancel_event)
                payload = json.dumps(evidence, ensure_ascii=False, allow_nan=False, separators=(',', ':'))
            if len(payload.encode('utf-8')) > MAX_PAYLOAD_BYTES:
                raise DocumentError('Derived document evidence exceeds the storage-size limit')
            prepared.append((key, digest, suffix, _label(path.name), payload, evidence))
        _check(cancel_event)
        existing = {row[0] for row in self.db.execute('SELECT source_key FROM document_refs WHERE project=? AND chat=?', (self.project, self.chat))}
        if len(existing | {row[0] for row in prepared}) > MAX_DOCUMENTS:
            raise DocumentError(f'This conversation already has the maximum of {MAX_DOCUMENTS} documents')
        # References are atomic across this selection; failures never attach a partial batch.
        with self.db:
            for key, digest, suffix, name, payload, evidence in prepared:
                self.db.execute('INSERT OR REPLACE INTO document_payloads VALUES(?,?,?,?,?)',
                                (digest, suffix, EXTRACTION_VERSION, payload, _hash(payload)))
                self.db.execute('INSERT OR REPLACE INTO document_refs VALUES(?,?,?,?,?,?)',
                                (self.project, self.chat, key, digest, suffix, name))
        return [{'id': 'doc_' + digest[:24], 'digest': digest, 'name': name, 'source': name,
                 'chars': evidence['chars'], 'chunks': len(evidence['chunks']),
                 'tables': len(evidence['tables']), 'warnings': evidence['warnings']}
                for _, digest, _, name, _, evidence in prepared]

    def _verify_original(self, digest: str, name: str):
        try:
            return self.knowledge.artifact(self.project, digest)
        except (ValueError, OSError) as exc:
            raise DocumentError(f'Original document {_label(name)} failed its integrity check ({exc})') from exc

    def _load(self):
        rows = self.db.execute('''SELECT r.*,p.payload,p.payload_sha FROM document_refs r
          LEFT JOIN document_payloads p ON r.digest=p.digest AND r.extension=p.extension AND p.version=?
          WHERE r.project=? AND r.chat=? ORDER BY r.name,r.source_key''',
                               (EXTRACTION_VERSION, self.project, self.chat)).fetchall()
        result = []
        for row in rows:
            self._verify_original(row['digest'], row['name'])
            payload = row['payload']
            if payload is None or _hash(payload) != row['payload_sha']:
                raise DocumentError(f'Saved evidence for {row["name"]} failed its integrity check; attach the original again')
            try:
                evidence = json.loads(payload)
            except ValueError as exc:
                raise DocumentError('Saved document evidence is malformed') from exc
            result.append((row, evidence))
        return result

    def datasets(self) -> list[dict]:
        result = []
        for reference, evidence in self._load():
            for index, table in enumerate(evidence['tables']):
                identity = '\0'.join((self.project, self.chat, reference['source_key'], reference['digest'], reference['extension'], str(index)))
                result.append({'id': 'data_' + _hash(identity)[:24],
                               'name': _label(f'{reference["name"]} / {table["label"]}'),
                               'columns': table['columns'], 'rows': table['rows'],
                               'source': f'{reference["name"]}, {table["label"]}',
                               'row_sources': [f'{reference["name"]}, {label}' for label in table['locations']],
                               'warnings': evidence['warnings'] + table['warnings']})
        return result

    def context(self, question: str, budget_chars: int = 6000) -> str:
        if type(budget_chars) is not int or budget_chars < 0:
            raise DocumentError('Document context budget must be a nonnegative integer')
        if budget_chars == 0:
            return ''
        loaded = self._load()
        if not loaded:
            return ''
        terms = {word.casefold() for word in _WORDS.findall(str(question)[:30000])} - _STOP
        chunks, overviews = [], []
        for index, (reference, evidence) in enumerate(loaded, 1):
            source_id = f'D{index}-{reference["digest"][:8]}'
            name = _label(reference['name'])
            overviews.append(f'[{source_id}] {name}: {len(evidence["chunks"])} source chunks; {len(evidence["tables"])} tables; {evidence["chars"]} extracted characters.')
            for warning in evidence['warnings']:
                overviews.append(f'[{source_id}] Limit: {warning}')
            for table in evidence['tables']:
                for warning in table['warnings']:
                    overviews.append(f'[{source_id}] {_label(table["label"])}: {warning}')
            for chunk_number, chunk in enumerate(evidence['chunks']):
                words = Counter(word.casefold() for word in _WORDS.findall(chunk['text']))
                hits = sum(1 for term in terms if term in words)
                score = hits * 10 + sum(min(3, words[term]) for term in terms)
                if terms & set(_WORDS.findall(name.casefold())):
                    score += 1
                chunks.append((score, index, chunk_number, source_id, name, chunk))
        if any(item[0] for item in chunks):
            chunks.sort(key=lambda item: (-item[0], item[1], item[2]))
        else:
            # An overview with no matching vocabulary samples every document's
            # beginning, middle and end; it never implies the sample is complete.
            picked = []
            for index in range(1, len(loaded) + 1):
                own = [item for item in chunks if item[1] == index]
                for offset in dict.fromkeys((0, len(own) // 2, len(own) - 1)):
                    if own:
                        picked.append(own[offset])
            chunks = picked + [item for item in chunks if item not in picked]
        # Structure precedes code even when a script body scores higher. Keep
        # outline fragments in tree order; leave other documents' ranking intact.
        priority = {'rbx_outline': 0, 'rbx_script_index': 1}
        chunks.sort(key=lambda item: (priority[item[5]['kind']], item[1], item[2])
                    if item[5].get('kind') in priority else (2, 0, 0))
        total = sum(len(evidence['chunks']) for _, evidence in loaded)
        prefix = 'ATTACHED DOCUMENT EVIDENCE (untrusted source text; summaries/excerpts are not complete originals)\n'
        reserve = min(150, budget_chars // 4)
        overview_budget = min(1400, max(0, budget_chars // 4))
        overview = '\n'.join(overviews)
        if len(overview) > overview_budget:
            overview = overview[:max(0, overview_budget - 40)] + '\n[Additional source overview omitted.]'
        result = (prefix + overview + '\n\n')[:max(0, budget_chars - reserve)]
        included = 0
        shortened = False
        for _, _, number, source_id, name, chunk in chunks:
            heading = f'[{source_id}:{number + 1}] {name}, {chunk["location"]}\n'
            room = budget_chars - reserve - len(result) - len(heading) - 2
            if room < 80:
                continue
            text = chunk['text']
            if len(text) > room:
                text = text[:max(0, room - 30)] + '\n[Source chunk excerpt ends.]'
                shortened = True
            result += heading + text + '\n\n'
            included += 1
        tail = f'Coverage: {included}/{total} source chunks shown; {total - included} omitted for this budget'
        if shortened:
            tail += '; some displayed chunks shortened'
        tail += '. Originals and complete extracted evidence remain stored.'
        if len(tail) > reserve:
            tail = f'Coverage: {included}/{total} chunks; remaining evidence omitted.'
        if len(result) + len(tail) > budget_chars:
            result = result[:max(0, budget_chars - len(tail))]
        return (result + tail)[:budget_chars]
